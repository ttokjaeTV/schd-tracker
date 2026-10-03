import json, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, LineChart, Reference

"""data/*.json 으로 엑셀 마스터(downloads/schd_dividend_tracker.xlsx)를 만든다.
수식은 엑셀·구글시트가 파일을 열 때 계산한다(fullCalcOnLoad)."""
from pathlib import Path
import calendar
ROOT=Path(__file__).resolve().parent.parent
J=lambda n: json.loads((ROOT/'data'/f'{n}.json').read_text(encoding='utf-8'))
META=J('meta'); KRJ=J('kr_etf'); HOLDJ=J('holdings')
def _sf(ex):
    f=1
    for sp in META['splits']:
        if ex<sp['date']: f*=sp['ratio']
    return f
rows=[dict(ex=d['ex'],rec=d['record'],pay=d['pay'],amt=d['amount'],split=_sf(d['ex']),close=d.get('close'),fx=d.get('fx'),src=d['source']) for d in J('dividends')]
LAST=rows[-1]; MAXY=int(LAST['ex'][:4]); LASTQ=(int(LAST['ex'][5:7])+2)//3
order=KRJ['order']; names={c:KRJ['etfs'][c]['name'] for c in order}
fee={c:KRJ['etfs'][c]['fee'] for c in order}; timing={c:KRJ['etfs'][c]['timing'] for c in order}
kr={c:[dict(gijunYmd=h[0],divAmt=h[1],divRt=h[2]) for h in KRJ['etfs'][c]['history']] for c in order}
_lk=max(KRJ['etfs'][c]['history'][-1][0] for c in order); _ky,_km=int(_lk[:4]),int(_lk[4:6])
KR_ASOF=dt.date(_ky,_km,calendar.monthrange(_ky,_km)[1])
FONT='맑은 고딕'
f_in=Font(name=FONT,size=10,color='0000FF')
f_fx=Font(name=FONT,size=10,color='000000')
f_link=Font(name=FONT,size=10,color='008000')
f_hd=Font(name=FONT,size=10,bold=True,color='FFFFFF')
f_t=Font(name=FONT,size=14,bold=True,color='1F2A44')
f_b=Font(name=FONT,size=10,bold=True)
f_n=Font(name=FONT,size=10)
f_s=Font(name=FONT,size=9,color='666666')
hd_fill=PatternFill('solid',fgColor='1F2A44')
in_fill=PatternFill('solid',fgColor='FFF9DB')
key_fill=PatternFill('solid',fgColor='FFFF00')
band=PatternFill('solid',fgColor='F4F6FA')
thin=Side(style='thin',color='D0D5DD'); bd=Border(bottom=thin)
C=Alignment(horizontal='center',vertical='center',wrap_text=True)
R=Alignment(horizontal='right',vertical='center')
USD4='$0.0000;-$0.0000;-'; PCT='0.00%;-0.00%;-'; KRW='#,##0.0"원";-#,##0.0"원";-'; KRW0='#,##0"원";-#,##0"원";-'
DATE='yyyy-mm-dd'

def header(ws,row,labels,widths=None):
    for i,t in enumerate(labels,1):
        c=ws.cell(row,i,t); c.font=f_hd; c.fill=hd_fill; c.alignment=C
    ws.row_dimensions[row].height=34
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w

wb=Workbook()
# ---------------- 사용법 ----------------
G=wb.active; G.title='사용법'
G.column_dimensions['A'].width=3; G.column_dimensions['B'].width=26; G.column_dimensions['C'].width=90
G['B2']='SCHD 배당 트래커 (마스터)'; G['B2'].font=f_t
G['B3']=f"최종 갱신: {META['updated']} · 반영 범위: {rows[0]['ex'][:4]}년 1분기 ~ {MAXY}년 {LASTQ}분기 (배당락 {LAST['ex']}) · GitHub Actions 자동 생성"; G['B3'].font=f_s
r=5
G.cell(r,2,'설정값').font=f_b; r+=1
G.cell(r,2,'미국 원천징수율').font=f_n
c=G.cell(r,3,META['withholding']); c.font=f_in; c.fill=key_fill; c.number_format='0.0%'; c.alignment=Alignment(horizontal='left')
c.comment=Comment('미국 상장 ETF 배당: 한국 거주자는 미국에서 15% 원천징수 (한미 조세조약). 국내 추가 원천징수 없음, 금융소득종합과세(연 2,000만원 초과) 합산 대상.','Claude')
WHT='사용법!$C$6'
r+=2
G.cell(r,2,'색상 규칙').font=f_b; r+=1
for k,v,font,fill in [('파란 글씨 + 연노랑 칸','직접 입력하는 값 (공시 배당금·날짜·종가·환율·분배금)',f_in,in_fill),('검정 글씨','수식 — 건드리지 마세요. 자동 계산됩니다',f_fx,None),('노란 칸','핵심 설정값',f_in,key_fill)]:
    a=G.cell(r,2,k); a.font=font
    if fill: a.fill=fill
    G.cell(r,3,v).font=f_n; r+=1
r+=1
G.cell(r,2,'분기 갱신 절차 (5분)').font=f_b; r+=1
steps=[
 ('① SCHD 배당 입력','[SCHD_배당] 시트 맨 아래 빈 줄(수식이 미리 깔린 줄)에 C~E, K, M, P 열만 입력합니다.\n  C 배당락일 · D 지급일 · E 공시 배당금 · K 배당락일 종가 · M 지급일 환율 · P 출처'),
 ('② 국내 SCHD형 분배 입력','[국내SCHD형_분배] 시트 맨 아래에 4종목의 월 분배금 3개월치를 붙여넣습니다 (종목코드·종목명·기준일·분배금·분배율).'),
 ('③ 비교 기준일 변경','[국내SCHD형_비교] B3 칸의 기준일을 분기 말일로 바꿉니다 (자동 생성 시 최신 분배월 말일로 맞춰짐).'),
 ('④ 결과 확인','[최신분기_요약] 시트가 자동으로 최신 분기 기준으로 바뀝니다. 카페 글 숫자는 여기서 가져가면 됩니다.'),
 ('⑤ 파일명','SCHD_배당트래커_YY년N분기_YYMMDD.xlsx 형식으로 저장'),
]
for k,v in steps:
    G.cell(r,2,k).font=f_n; c=G.cell(r,3,v); c.font=f_n; c.alignment=Alignment(wrap_text=True,vertical='top'); G.row_dimensions[r].height=32 if '\n' in v else 18; r+=1
r+=1
G.cell(r,2,'데이터 출처').font=f_b; r+=1
src=[
 ('SCHD 배당 (2016.12~)','Schwab Asset Management 공식 분배 내역 CSV — schwabassetmanagement.com/products/schd (Distributions → Export Data)'),
 ('SCHD 배당 (2012~2016.09)','dividendhistory.net 원배당(Unadjusted) — Schwab CSV가 2016.12부터만 제공. Yahoo Finance 분할조정값과 교차 확인'),
 ('배당락일 종가','Yahoo Finance 일별 종가 (2024.10.11 3:1 분할 반영, 배당 미조정)'),
 ('지급일 환율','하나은행 고시 매매기준율 (다음금융 일별 환율, 지급일 당일 최종 회차)'),
 ('국내 SCHD형 분배금','FunETF 분배금 이력 (기준일·분배금·분배율) — 운용사 공시와 같은 원천'),
 ('총보수',KRJ['fee_source']),
]
for k,v in src:
    G.cell(r,2,k).font=f_n; c=G.cell(r,3,v); c.font=f_n; c.alignment=Alignment(wrap_text=True); r+=1
r+=1
G.cell(r,2,'알아둘 점').font=f_b; r+=1
notes=[
 '· 2024-10-11 SCHD 3:1 주식분할. 그 이전 배당은 공시 금액 ÷ 3 으로 환산해 지금 1주 기준으로 맞춥니다 (F열 자동).',
 '· 연간 성장률은 4회 배당이 모두 확정된 해만 계산합니다. 진행 중인 해는 [연도별] 시트의 \'올해 누적 vs 작년 동기\'로 봅니다.',
 '· \'평균 성장률(산술)\'과 \'CAGR(연평균 복리)\'은 다른 숫자입니다. 글에 쓸 때 용어를 구분하세요.',
 '· 진행 중인 해의 남은 분기 \'예상\' 값은 그해 확정 분기의 단순 평균입니다. 과거 14년 중 13년은 4분기가 1~3분기 평균보다 높았으니(보통 +8~16%) 연간 예상치는 보수적(낮게) 나올 가능성이 큽니다.',
 '· 이 파일은 GitHub에서 새 배당이 반영될 때마다 다시 만들어집니다(차트 범위도 자동). 손으로 고친 내용은 다음 생성 때 덮어써지니, 고칠 값은 data 폴더 JSON에서 고치세요.',
 '· 국내 SCHD형 ETF 분배율은 매월 기준일 가격 대비 분배율을 12개월 합산한 값입니다(근사치).',
]
for n in notes:
    c=G.cell(r,2,n); c.font=f_n; G.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); r+=1

# ---------------- SCHD_배당 ----------------
S=wb.create_sheet('SCHD_배당')
heads=['연도','분기','배당락일','지급일','공시 배당금\n(USD, 당시 1주)','분할\n환산','수정 배당금\n(USD, 현재 1주)','전분기\n대비','전년 동기\n대비','최근 4회\n합계(TTM)','배당락일\n종가(USD)','TTM\n배당수익률','지급일\n환율(원)','1주 배당\n원화(세전)','1주 배당\n원화(미국15%후)','출처']
header(S,1,heads,[7,6,11.5,11.5,13,6,13,9,9,11,11,10,10,11,12,22])
S.freeze_panes='C2'
N=len(rows); FUT=30; last=1+N+FUT
for i in range(N+FUT):
    rr=i+2; d=rows[i] if i<N else None
    S.cell(rr,1,f'=IF($C{rr}="","",YEAR($C{rr}))').number_format='0'
    S.cell(rr,2,f'=IF($C{rr}="","",ROUNDUP(MONTH($C{rr})/3,0))').number_format='"Q"0'
    for col in (3,4,5,11,13,16):
        c=S.cell(rr,col); c.font=f_in; c.fill=in_fill
    if d:
        S.cell(rr,3,dt.date.fromisoformat(d['ex'])); S.cell(rr,4,dt.date.fromisoformat(d['pay']))
        S.cell(rr,5,d['amt']); S.cell(rr,11,d['close']); S.cell(rr,13,d['fx']); S.cell(rr,16,d['src'])
    S.cell(rr,3).number_format=DATE; S.cell(rr,4).number_format=DATE
    S.cell(rr,5).number_format='$0.0000'; S.cell(rr,11).number_format='$0.00'; S.cell(rr,13).number_format='#,##0.0'
    S.cell(rr,6,f'=IF($C{rr}="","",IF($C{rr}<DATE(2024,10,11),3,1))').number_format='0'
    S.cell(rr,7,f'=IF($C{rr}="","",E{rr}/F{rr})').number_format=USD4
    S.cell(rr,8,'' if rr==2 else f'=IF(OR($C{rr}="",G{rr-1}=""),"",G{rr}/G{rr-1}-1)').number_format=PCT
    S.cell(rr,9,'' if rr<6 else f'=IF(OR($C{rr}="",G{rr-4}=""),"",G{rr}/G{rr-4}-1)').number_format=PCT
    S.cell(rr,10,'' if rr<5 else f'=IF($C{rr}="","",SUM(G{rr-3}:G{rr}))').number_format=USD4
    S.cell(rr,12,f'=IF(OR(J{rr}="",K{rr}=""),"",J{rr}/K{rr})').number_format=PCT
    S.cell(rr,14,f'=IF(OR(G{rr}="",M{rr}=""),"",G{rr}*M{rr})').number_format=KRW
    S.cell(rr,15,f'=IF(N{rr}="","",N{rr}*(1-{WHT}))').number_format=KRW
    for col in range(1,17):
        c=S.cell(rr,col)
        if col not in (3,4,5,11,13,16): c.font=f_fx
        c.border=bd
        if col in (1,2,6): c.alignment=Alignment(horizontal='center')
S['F1'].comment=Comment('2024-10-11 3:1 분할 이전 배당은 ÷3. 앞으로 분할이 또 생기면 이 열의 날짜 조건을 수정하세요.','Claude')
S['E1'].comment=Comment('운용사(Schwab) 공시 금액 그대로 입력. 분할 전 배당은 분할 전 1주 기준 금액입니다.','Claude')
S['M1'].comment=Comment('지급일 하나은행 매매기준율(최종 회차). 실제 입금 환율은 증권사·지급 시점에 따라 다릅니다.','Claude')
S['O1'].comment=Comment('미국 원천징수율은 [사용법] C6 칸에서 변경','Claude')
MS="SCHD_배당"; RNG=lambda col: f"{MS}!${col}$2:${col}${last}"

# ---------------- 연도별 ----------------
Y=wb.create_sheet('연도별')
header(Y,1,['연도','배당\n횟수','연간 배당\n(USD, 현재 1주)','상태','전년 대비\n성장률','연속 증가\n(년)','올해 누적\nvs 작년 동기','연간 배당\n원화(세전)','미지급 분기\n예상분','연간 배당\n(예상 포함)','예상\n성장률'],[8,8,15,13,11,10,13,14,12,13,10])
Y.freeze_panes='A2'
years=list(range(2012,MAXY+6))
for i,yv in enumerate(years):
    rr=i+2
    Y.cell(rr,1,yv).font=f_fx; Y.cell(rr,1).alignment=Alignment(horizontal='center')
    Y.cell(rr,2,f'=COUNTIFS({RNG("A")},A{rr})')
    Y.cell(rr,3,f'=IF(B{rr}=0,"",SUMIFS({RNG("G")},{RNG("A")},A{rr}))').number_format=USD4
    Y.cell(rr,4,f'=IF(B{rr}=0,"",IF(B{rr}>=4,"확정","진행중("&B{rr}&"회)"))')
    Y.cell(rr,5,'' if rr==2 else f'=IF(AND(B{rr}>=4,B{rr-1}>=4),C{rr}/C{rr-1}-1,"")').number_format=PCT
    Y.cell(rr,6,'' if rr==2 else f'=IF(E{rr}="","",IF(E{rr}>0,IF(F{rr-1}="",1,F{rr-1}+1),0))')
    Y.cell(rr,7,'' if rr==2 else f'=IF(AND(B{rr}>0,B{rr}<4),C{rr}/SUMIFS({RNG("G")},{RNG("A")},A{rr-1},{RNG("B")},"<="&B{rr})-1,"")').number_format=PCT
    Y.cell(rr,8,f'=IF(B{rr}=0,"",SUMIFS({RNG("N")},{RNG("A")},A{rr}))').number_format=KRW
    Y.cell(rr,9,f'=IF(AND(B{rr}>0,B{rr}<4),(4-B{rr})*C{rr}/B{rr},"")').number_format=USD4
    Y.cell(rr,10,f'=IF(B{rr}=0,"",C{rr}+IF(I{rr}="",0,I{rr}))').number_format=USD4
    Y.cell(rr,11,'' if rr==2 else f'=IF(AND(I{rr}<>"",B{rr-1}>=4),J{rr}/C{rr-1}-1,"")').number_format=PCT
    for col in (9,10,11): Y.cell(rr,col).fill=PatternFill('solid',fgColor='FBEFDF')
    for col in range(1,12):
        c=Y.cell(rr,col); c.border=bd
        if col>1: c.font=f_fx
        if col in (2,4,6): c.alignment=Alignment(horizontal='center')
Yl=len(years)+1
Y.cell(2,6,'')  # 2012 기준점
Y['F1'].comment=Comment('2012년(첫 전체 연도)부터 전년 대비 증가가 이어진 햇수','Claude')
Y['H1'].comment=Comment('각 분기 지급일 환율로 환산한 세전 원화 합계','Claude')
# stats block
b=Yl+2
Y.cell(b,1,'핵심 지표 (최근 확정연도 기준, 자동)').font=f_b
YR=lambda col: f'$'+col+'$2:$'+col+f'${Yl}'
items=[
 ('최근 확정연도', f'=_xlfn.MAXIFS({YR("A")},{YR("B")},">=4")','0'),
 ('그해 연간 배당', f'=INDEX({YR("C")},MATCH(B{b+1},{YR("A")},0))',USD4),
 ('그해 성장률', f'=INDEX({YR("E")},MATCH(B{b+1},{YR("A")},0))',PCT),
 ('5년 CAGR', f'=(B{b+2}/INDEX({YR("C")},MATCH(B{b+1}-5,{YR("A")},0)))^(1/5)-1',PCT),
 ('10년 CAGR', f'=(B{b+2}/INDEX({YR("C")},MATCH(B{b+1}-10,{YR("A")},0)))^(1/10)-1',PCT),
 ('10년 평균 성장률(산술)', f'=AVERAGEIFS({YR("E")},{YR("A")},">"&(B{b+1}-10),{YR("A")},"<="&B{b+1})',PCT),
 ('연속 증가 연수', f'=INDEX({YR("F")},MATCH(B{b+1},{YR("A")},0))','0"년"'),
 ('진행 중인 해', f'=MAX({RNG("A")})','0'),
 ('  예상 연간 배당 (남은 분기 = 올해 확정 분기 평균)', f'=IF(INDEX({YR("I")},MATCH(B{b+8},{YR("A")},0))="","",INDEX({YR("J")},MATCH(B{b+8},{YR("A")},0)))',USD4),
 ('  예상 성장률 (남은 분기 = 올해 확정 분기 평균)', f'=IF(B{b+9}="","",B{b+9}/INDEX({YR("C")},MATCH(B{b+8}-1,{YR("A")},0))-1)',PCT),
 ('  예상 연간 배당 (남은 분기 = 확정 분기 평균 ×110%)', f'=IF(B{b+9}="","",INDEX({YR("C")},MATCH(B{b+8},{YR("A")},0))+INDEX({YR("I")},MATCH(B{b+8},{YR("A")},0))*1.1)',USD4),
 ('  예상 성장률 (남은 분기 = 확정 분기 평균 ×110%)', f'=IF(B{b+11}="","",B{b+11}/INDEX({YR("C")},MATCH(B{b+8}-1,{YR("A")},0))-1)',PCT),
]
for j,(k,fml,fmt) in enumerate(items,1):
    Y.cell(b+j,1,k).font=f_n
    c=Y.cell(b+j,2,fml); c.number_format=fmt; c.font=f_fx
Y.column_dimensions['A'].width=40
Y.cell(b+len(items)+1,1,'※ CAGR = (최근 확정연도 배당 ÷ N년 전 배당)^(1/N) − 1. 산술평균과 다른 숫자입니다.').font=f_s
Y.cell(b+len(items)+2,1,'※ 예상치: 진행 중인 해의 남은 분기를 그해 확정 분기 평균으로 채운 단순 계산. 과거엔 4분기가 1·2·3분기 평균보다 높았던 해가 대부분이라, 남은 분기가 확정 분기 평균보다 10% 높다고 가정한 값도 함께 표시').font=f_s

Y['I1'].comment=Comment('진행 중인 해의 미지급 분기를, 그해 이미 확정된 분기 배당의 단순 평균으로 채운 값 (예측치)','Claude')
# chart helper (M~Q): 차트용 시리즈 — 확정/예상 분리
for j,t in enumerate(['차트용\n연도','확정\n배당','예상\n추가분','확정\n성장률','예상\n성장률','전체 평균\n배당','전체 평균\n성장률','연간 합계\n(예상 포함)'],13):
    c=Y.cell(1,j,t); c.font=f_hd; c.fill=PatternFill('solid',fgColor='5E6E78'); c.alignment=C
    Y.column_dimensions[L(j)].width=9
for i,yv in enumerate(years):
    rr=i+2
    Y.cell(rr,13,f'=IF(B{rr}=0,"",TEXT(A{rr},"0"))')
    Y.cell(rr,14,f'=IF(B{rr}=0,"",C{rr})').number_format=USD4
    Y.cell(rr,15,f'=IF(I{rr}="","",I{rr})').number_format=USD4
    Y.cell(rr,16,f'=IF(E{rr}="","",E{rr})').number_format=PCT
    Y.cell(rr,17,f'=IF(K{rr}="","",K{rr})').number_format=PCT
    Y.cell(rr,18,f'=IF(B{rr}=0,"",AVERAGEIFS({YR("C")},{YR("B")},">=4"))').number_format=USD4
    Y.cell(rr,19,f'=IF(B{rr}=0,"",AVERAGE({YR("E")}))').number_format=PCT
    Y.cell(rr,20,f'=IF(B{rr}=0,"",J{rr})').number_format=USD4
    for col in range(13,21): Y.cell(rr,col).font=f_s
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
CH_END=MAXY-2012+2  # 데이터가 있는 마지막 연도 행
from openpyxl.chart.label import DataLabelList
def avg_line(ref,cats):
    lc=LineChart(); lc.add_data(ref,titles_from_data=True); lc.set_categories(cats)
    sr=lc.series[0]; sr.marker.symbol='none'; sr.smooth=False
    sr.graphicalProperties=GraphicalProperties(); sr.graphicalProperties.line=LineProperties(solidFill='2F5D8A',w=22000,prstDash='dash')
    return lc
def label_line(ref,cats,fmt):
    lc=LineChart(); lc.add_data(ref,titles_from_data=True); lc.set_categories(cats)
    sr=lc.series[0]; sr.marker.symbol='none'
    sr.graphicalProperties=GraphicalProperties(); sr.graphicalProperties.line=LineProperties(noFill=True)
    sr.dLbls=DataLabelList(); sr.dLbls.showVal=True; sr.dLbls.position='t'; sr.dLbls.numFmt=fmt
    for k in ('showSerName','showCatName','showLegendKey','showPercent'): setattr(sr.dLbls,k,False)
    return lc
def solid(series,hexc,dash=False):
    series.graphicalProperties=GraphicalProperties(solidFill=hexc)
    series.graphicalProperties.line=LineProperties(solidFill='A8661B' if dash else hexc, prstDash='dash' if dash else None)
ch=BarChart(); ch.type='col'; ch.grouping='stacked'; ch.overlap=100
ch.title='SCHD 연간 배당 (USD) · 진한색 확정 / 연한색 예상'; ch.height=9; ch.width=20
ch.add_data(Reference(Y,min_col=14,min_row=1,max_row=CH_END),titles_from_data=True)
ch.add_data(Reference(Y,min_col=15,min_row=1,max_row=CH_END),titles_from_data=True)
ch.set_categories(Reference(Y,min_col=13,min_row=2,max_row=CH_END))
solid(ch.series[0],'C98A3A'); solid(ch.series[1],'F1D9B5',dash=True)
ch.y_axis.numFmt='$0.00'; ch.y_axis.majorGridlines=None; ch.legend.position='b'
ch.y_axis.delete=False; ch.x_axis.delete=False
_cats=Reference(Y,min_col=13,min_row=2,max_row=CH_END)
ch+=avg_line(Reference(Y,min_col=18,min_row=1,max_row=CH_END),_cats)
ch+=label_line(Reference(Y,min_col=20,min_row=1,max_row=CH_END),_cats,'$0.0000')
ch.title='SCHD 연간 배당 (USD) · 진한색 확정 / 연한색 예상 / 점선 전체 평균'
Y.add_chart(ch,'V2')
g=BarChart(); g.type='col'; g.grouping='clustered'; g.overlap=100
g.title='SCHD 연간 배당 성장률 · 진한색 확정 / 연한색 예상'; g.height=9; g.width=20
g.add_data(Reference(Y,min_col=16,min_row=1,max_row=CH_END),titles_from_data=True)
g.add_data(Reference(Y,min_col=17,min_row=1,max_row=CH_END),titles_from_data=True)
g.set_categories(Reference(Y,min_col=13,min_row=2,max_row=CH_END))
solid(g.series[0],'17784A'); solid(g.series[1],'F1D9B5',dash=True)
g.y_axis.numFmt='0%'; g.legend.position='b'; g.y_axis.delete=False; g.x_axis.delete=False
g+=avg_line(Reference(Y,min_col=19,min_row=1,max_row=CH_END),_cats)
g.title='SCHD 연간 배당 성장률 · 진한색 확정 / 연한색 예상 / 점선 전체 평균'
Y.add_chart(g,'V22')

# ---------------- 분기표 ----------------
Q=wb.create_sheet('분기표')
header(Q,1,['연도','Q1','Q2','Q3','Q4','연간','','연도','Q1\n전년비','Q2\n전년비','Q3\n전년비','Q4\n전년비'],[8,10,10,10,10,10,3,8,9,9,9,9])
Q.cell(1,7).fill=PatternFill(None)
for i,yv in enumerate(years):
    rr=i+2
    Q.cell(rr,1,yv); Q.cell(rr,8,yv)
    for q in range(1,5):
        Q.cell(rr,1+q,f'=IF(COUNTIFS({RNG("A")},$A{rr},{RNG("B")},{q})=0,"",SUMIFS({RNG("G")},{RNG("A")},$A{rr},{RNG("B")},{q}))').number_format=USD4
        if rr>2:
            col=L(1+q)
            Q.cell(rr,8+q,f'=IF(OR({col}{rr}="",{col}{rr-1}=""),"",{col}{rr}/{col}{rr-1}-1)').number_format=PCT
    Q.cell(rr,6,f'=IF(SUM(B{rr}:E{rr})=0,"",SUM(B{rr}:E{rr}))').number_format=USD4
    for col in range(1,13):
        if col==7: continue
        c=Q.cell(rr,col); c.font=f_fx; c.border=bd
        if col in (1,8): c.alignment=Alignment(horizontal='center')
Q.freeze_panes='B2'

# ---------------- 국내SCHD형_분배 ----------------
K=wb.create_sheet('국내SCHD형_분배')
header(K,1,['종목코드','종목명','기준일','분배금(원)','분배율','비고'],[10,24,12,11,9,30])
K.freeze_panes='A2'
recs=[]
for code in order:
    for x in sorted(kr[code],key=lambda z:z['gijunYmd']):
        recs.append((code,names[code],dt.datetime.strptime(x['gijunYmd'],'%Y%m%d').date(),x['divAmt'],x['divRt']/100))
for i,(code,nm,d,amt,rt) in enumerate(recs):
    rr=i+2
    for col,v,fmt in [(1,code,'@'),(2,nm,None),(3,d,DATE),(4,amt,'#,##0'),(5,rt,'0.00%')]:
        c=K.cell(rr,col,v); c.font=f_in; c.fill=in_fill; c.border=bd
        if fmt: c.number_format=fmt
KL=len(recs)+1+200
for rr in range(len(recs)+2,KL+1):
    for col in range(1,6):
        c=K.cell(rr,col); c.fill=in_fill; c.font=f_in; c.border=bd
    K.cell(rr,1).number_format='@'; K.cell(rr,3).number_format=DATE; K.cell(rr,4).number_format='#,##0'; K.cell(rr,5).number_format='0.00%'
K['C1'].comment=Comment('분배금 지급 기준일(권리 확정일). 출처: FunETF 분배 이력','Claude')
K['E1'].comment=Comment('기준일 가격 대비 1회 분배율','Claude')
KR=lambda col: f"국내SCHD형_분배!${col}$2:${col}${KL}"

# ---------------- 국내SCHD형_비교 ----------------
V=wb.create_sheet('국내SCHD형_비교')
V['A1']='국내상장 SCHD형 ETF vs SCHD'; V['A1'].font=f_t
V['A3']='비교 기준일'; V['A3'].font=f_b
V['B3']=KR_ASOF; V['B3'].font=f_in; V['B3'].fill=key_fill; V['B3'].number_format=DATE
V['C3']='← 분기마다 분기 말일로 바꾸세요. 아래 표가 이 날짜 기준 최근 12개월로 다시 계산됩니다.'; V['C3'].font=f_s
hrow=5
header(V,hrow,['종목','코드','총보수\n(연)','분배\n시점','최근 분배일','최근 분배금','12개월\n분배 횟수','12개월\n분배금 합계','12개월\n분배율 합계'],[24,9,9,10,12,11,9,12,11])
V.row_dimensions[hrow].height=34
for j,code in enumerate(order):
    rr=hrow+1+j
    V.cell(rr,1,names[code]).font=f_n
    c=V.cell(rr,2,code); c.number_format='@'; c.font=f_n
    c=V.cell(rr,3,fee[code]); c.font=f_in; c.fill=in_fill; c.number_format='0.0000%'
    c=V.cell(rr,4,timing[code]); c.font=f_in; c.fill=in_fill
    V.cell(rr,5,f'=_xlfn.MAXIFS({KR("C")},{KR("A")},$B{rr},{KR("C")},"<="&$B$3)').number_format=DATE
    V.cell(rr,6,f'=SUMIFS({KR("D")},{KR("A")},$B{rr},{KR("C")},E{rr})').number_format=KRW0
    V.cell(rr,7,f'=COUNTIFS({KR("A")},$B{rr},{KR("C")},">"&EDATE($B$3,-12),{KR("C")},"<="&$B$3)')
    V.cell(rr,8,f'=SUMIFS({KR("D")},{KR("A")},$B{rr},{KR("C")},">"&EDATE($B$3,-12),{KR("C")},"<="&$B$3)').number_format=KRW0
    V.cell(rr,9,f'=SUMIFS({KR("E")},{KR("A")},$B{rr},{KR("C")},">"&EDATE($B$3,-12),{KR("C")},"<="&$B$3)').number_format=PCT
    for col in range(1,10):
        V.cell(rr,col).border=bd
        if col>=5: V.cell(rr,col).font=f_fx
sr=hrow+5
V.cell(sr,1,'SCHD (미국 직투, 참고)').font=f_b
V.cell(sr,2,'SCHD')
c=V.cell(sr,3,META['schd_fee']); c.font=f_in; c.fill=in_fill; c.number_format='0.00%'
V.cell(sr,4,'분기(3·6·9·12월)')
V.cell(sr,5,f'=_xlfn.MAXIFS({RNG("C")},{RNG("C")},"<="&$B$3)').number_format=DATE
V.cell(sr,6,f'=INDEX({RNG("G")},MATCH(E{sr},{RNG("C")},0))').number_format=USD4
V.cell(sr,7,f'=COUNTIFS({RNG("C")},">"&EDATE($B$3,-12),{RNG("C")},"<="&$B$3)')
V.cell(sr,8,f'=INDEX({RNG("J")},MATCH(E{sr},{RNG("C")},0))').number_format=USD4
V.cell(sr,9,f'=INDEX({RNG("L")},MATCH(E{sr},{RNG("C")},0))').number_format=PCT
for col in range(1,10): V.cell(sr,col).border=bd
V.cell(sr+1,1,'※ SCHD 행: 금액은 USD, 수익률은 최근 4회 배당 ÷ 배당락일 종가. 총보수 0.06%는 Schwab 공시.').font=f_s
V.cell(sr+2,1,f"※ 국내 ETF 총보수: {KRJ['fee_source']}. 실부담비용(기타비용·매매중개수수료 포함)은 이보다 높습니다.").font=f_s
V.cell(sr+3,1,'※ 국내상장 해외주식형 ETF 분배금은 배당소득세 15.4% 과세 (연금저축·IRP·ISA에서는 과세이연/비과세 혜택).').font=f_s

# quarterly table
qr=sr+6
V.cell(qr-1,1,'분기별 분배금 합계 (기준일 기준 분기 귀속)').font=f_b
header_cells=['분기','시작일','종료일']+[names[c].split()[0] for c in order]+['SCHD\n(USD)','SCHD 원화\n(세전)']
for i,t in enumerate(header_cells,1):
    c=V.cell(qr,i,t); c.font=f_hd; c.fill=hd_fill; c.alignment=C
V.row_dimensions[qr].height=30
quarters=[(y,q) for y in range(2022,MAXY+3) for q in range(1,5)]
for i,(yv,q) in enumerate(quarters):
    rr=qr+1+i
    V.cell(rr,1,f'{yv} Q{q}').alignment=Alignment(horizontal='center')
    V.cell(rr,2,dt.date(yv,3*q-2,1)).number_format=DATE
    V.cell(rr,3,f'=EDATE(B{rr},3)-1').number_format=DATE
    for j,code in enumerate(order):
        V.cell(rr,4+j,f'=IF(COUNTIFS({KR("A")},"{code}",{KR("C")},">="&$B{rr},{KR("C")},"<="&$C{rr})=0,"",SUMIFS({KR("D")},{KR("A")},"{code}",{KR("C")},">="&$B{rr},{KR("C")},"<="&$C{rr}))').number_format=KRW0
    V.cell(rr,8,f'=IF(COUNTIFS({RNG("C")},">="&$B{rr},{RNG("C")},"<="&$C{rr})=0,"",SUMIFS({RNG("G")},{RNG("C")},">="&$B{rr},{RNG("C")},"<="&$C{rr}))').number_format=USD4
    V.cell(rr,9,f'=IF(H{rr}="","",SUMIFS({RNG("N")},{RNG("C")},">="&$B{rr},{RNG("C")},"<="&$C{rr}))').number_format=KRW
    for col in range(1,10):
        c=V.cell(rr,col); c.border=bd
        if col>1: c.font=f_fx
V.column_dimensions['A'].width=24
V.freeze_panes='B6'

# ---------------- 최신분기_요약 ----------------
Z=wb.create_sheet('최신분기_요약')
Z.column_dimensions['A'].width=34; Z.column_dimensions['B'].width=16; Z.column_dimensions['C'].width=50
Z['A1']='SCHD 최신 분기 요약 (카페 글용 숫자)'; Z['A1'].font=f_t
Z['A2']='[SCHD_배당]에 새 분기를 입력하면 자동으로 바뀝니다.'; Z['A2'].font=f_s
LR=f'(COUNT({RNG("C")})+1)'
IX=lambda col: f'=INDEX({MS}!${col}:${col},{LR})'
Yn=lambda col,yr: f'INDEX(연도별!${col}$2:${col}${Yl},MATCH({yr},연도별!$A$2:$A${Yl},0))'
lines=[
 ('최신 배당', None, None, None),
 ('연도', IX('A'),'0',None),
 ('분기', IX('B'),'"Q"0',None),
 ('배당락일', IX('C'),DATE,None),
 ('지급일', IX('D'),DATE,None),
 ('배당금 (USD)', IX('G'),USD4,None),
 ('전분기 대비', IX('H'),'+0.00%;-0.00%;0.00%',None),
 ('전년 동기 대비', IX('I'),'+0.00%;-0.00%;0.00%',None),
 ('전년 동기 배당금 (USD)', f'=INDEX({MS}!$G:$G,{LR}-4)',USD4,None),
 ('최근 4회 합계 TTM (USD)', IX('J'),USD4,None),
 ('배당락일 종가 (USD)', IX('K'),'$0.00',None),
 ('TTM 배당수익률', IX('L'),PCT,None),
 ('지급일 환율', IX('M'),'#,##0.0"원"',None),
 ('1주 배당 원화 (세전)', IX('N'),KRW,None),
 ('1주 배당 원화 (미국 15% 원천징수 후)', IX('O'),KRW,None),
 ('', None,None,None),
 ('올해 흐름', None,None,None),
 ('올해 누적 배당 (USD)', f'={Yn("C","B4")}',USD4,None),
 ('작년 같은 기간 누적 (USD)', f'=SUMIFS({RNG("G")},{RNG("A")},B4-1,{RNG("B")},"<="&B5)',USD4,None),
 ('올해 누적 vs 작년 동기', f'=B20/B21-1','+0.00%;-0.00%;0.00%',None),
 ('', None,None,None),
 ('장기 성장', None,None,None),
 ('최근 확정연도', f'=_xlfn.MAXIFS(연도별!$A$2:$A${Yl},연도별!$B$2:$B${Yl},">=4")','0',None),
 ('그해 연간 배당 (USD)', f'={Yn("C","B25")}',USD4,None),
 ('그해 성장률', f'={Yn("E","B25")}',PCT,None),
 ('5년 CAGR', f'=(B26/{Yn("C","B25-5")})^(1/5)-1',PCT,None),
 ('10년 CAGR', f'=(B26/{Yn("C","B25-10")})^(1/10)-1',PCT,None),
 ('10년 평균 성장률 (산술)', f'=AVERAGEIFS(연도별!$E$2:$E${Yl},연도별!$A$2:$A${Yl},">"&(B25-10),연도별!$A$2:$A${Yl},"<="&B25)',PCT,None),
 ('연간 배당 연속 증가 (2012년 이후 집계)', f'={Yn("F","B25")}','0"년 연속"',None),
]
r=3
for k,fml,fmt,_ in lines:
    if fml is None:
        if k: Z.cell(r,1,k).font=Font(name=FONT,size=11,bold=True,color='1F2A44')
    else:
        Z.cell(r,1,k).font=f_n
        c=Z.cell(r,2,fml); c.number_format=fmt; c.font=f_fx; c.alignment=R; c.border=bd
        Z.cell(r,1).border=bd
    r+=1
# sanity: verify row anchors
assert Z['A4'].value=='연도' and Z['A5'].value=='분기' and Z['A20'].value=='올해 누적 배당 (USD)' and Z['A21'].value.startswith('작년') and Z['A25'].value=='최근 확정연도' and Z['A26'].value.startswith('그해 연간')
Z['C9'].value='+면 인상, −면 인하'; Z['C9'].font=f_s
Z['C22'].value='배당 횟수가 같은 기간끼리 비교'; Z['C22'].font=f_s

# quarterly line chart on SCHD sheet
Q2=wb.create_sheet('차트데이터')
header(Q2,1,['연도','분기','라벨','확정 배당','예상 배당','전체 평균'],[8,6,9,11,11,11])
Q2['G1']='분기 차트용 표. 예상 배당 = 올해(데이터가 있는 마지막 연도) 미지급 분기를 그해 확정 분기의 단순 평균으로 채운 값'; Q2['G1'].font=f_s
CY=f'MAX({RNG("A")})'
qrow=2
for yv in range(2012,MAXY+1):
    for q in range(1,5):
        rr=qrow
        Q2.cell(rr,1,yv); Q2.cell(rr,2,q).number_format='"Q"0'
        Q2.cell(rr,3,f"{str(yv)[2:]} Q{q}")
        Q2.cell(rr,4,f'=IF(COUNTIFS({RNG("A")},A{rr},{RNG("B")},B{rr})=0,"",SUMIFS({RNG("G")},{RNG("A")},A{rr},{RNG("B")},B{rr}))').number_format=USD4
        Q2.cell(rr,5,f'=IF(AND(D{rr}="",A{rr}={CY},COUNTIFS({RNG("A")},A{rr})>0),SUMIFS({RNG("G")},{RNG("A")},A{rr})/COUNTIFS({RNG("A")},A{rr}),"")').number_format=USD4
        Q2.cell(rr,6,f'=AVERAGE($D$2:$D${(MAXY-2012+1)*4+1})').number_format=USD4
        for col in range(1,7): Q2.cell(rr,col).font=f_fx; Q2.cell(rr,col).border=bd
        qrow+=1
qc=BarChart(); qc.type='col'; qc.grouping='clustered'; qc.overlap=100; qc.gapWidth=40
qc.title='SCHD 분기 배당 (USD, 현재 1주 기준) · 진한색 확정 / 연한색 예상'; qc.height=9; qc.width=30
qc.add_data(Reference(Q2,min_col=4,min_row=1,max_row=qrow-1),titles_from_data=True)
qc.add_data(Reference(Q2,min_col=5,min_row=1,max_row=qrow-1),titles_from_data=True)
qc.set_categories(Reference(Q2,min_col=3,min_row=2,max_row=qrow-1))
solid(qc.series[0],'C98A3A'); solid(qc.series[1],'F1D9B5',dash=True)
qc.y_axis.numFmt='$0.00'; qc.legend.position='b'; qc.y_axis.delete=False; qc.x_axis.delete=False
qc.x_axis.tickLblSkip=4
qc+=avg_line(Reference(Q2,min_col=6,min_row=1,max_row=qrow-1),Reference(Q2,min_col=3,min_row=2,max_row=qrow-1))
qc.title='SCHD 분기 배당 (USD) · 진한색 확정 / 연한색 예상 / 점선 전체 평균'
S.add_chart(qc,'R2')


# ---------------- 구성종목 ----------------
import pandas as _pd
from openpyxl.chart import DoughnutChart
from openpyxl.chart.series import DataPoint
H=wb.create_sheet('구성종목')
hold=_pd.DataFrame([dict(Symbol=i['symbol'],Percent=i['weight'],Name=i['name'],Sector=i['sector']) for i in HOLDJ['items']]).fillna('')
KN=META['names_ko']
SK=[('Consumer Staples','필수소비재'),('Health Care','헬스케어'),('Energy','에너지'),('Information Technology','정보기술'),('Industrials','산업재'),('Financials','금융'),('Consumer Discretionary','경기소비재'),('Communication Services','커뮤니케이션'),('Utilities','유틸리티'),('Real Estate','부동산'),('Materials','소재')]
H['A1']='SCHD 구성종목'; H['A1'].font=f_t
H['A2']='기준일'; H['A2'].font=f_b
H['B2']=dt.date.fromisoformat(HOLDJ['asOf']); H['B2'].font=f_in; H['B2'].fill=key_fill; H['B2'].number_format=DATE
H['C2']='← Schwab 공식 보유종목 CSV 기준일. 자동 갱신 시 매주 최신 목록으로 바뀝니다.'; H['C2'].font=f_s
# 전체 목록 (H~L)
hs=4
for j,tl in enumerate(['티커','비중','종목명(영문)','섹터(GICS)','한글명','섹터(한글)'],8):
    c=H.cell(hs,j,tl); c.font=f_hd; c.fill=hd_fill; c.alignment=C
for j,w in zip(range(8,14),[8,9,30,22,18,12]): H.column_dimensions[L(j)].width=w
HL=hs+len(hold)+60
for i,row in hold.iterrows():
    rr=hs+1+i
    for col,v,fmt in [(8,row.Symbol,None),(9,float(row.Percent)/100,'0.00%'),(10,row.Name,None),(11,row.Sector,None),(12,KN.get(row.Symbol,''),None)]:
        c=H.cell(rr,col,v); c.font=f_in; c.fill=in_fill
        if fmt: c.number_format=fmt
for rr in range(hs+1,HL+1):
    for col in range(8,13):
        c=H.cell(rr,col); c.fill=in_fill; c.font=f_in
    H.cell(rr,9).number_format='0.00%'
    H.cell(rr,13,f'=IF(H{rr}="","",IFERROR(VLOOKUP(K{rr},$P$5:$Q$15,2,FALSE),"현금·선물"))').font=f_fx
# 섹터 매핑 표 (P~Q)
H['P4']='GICS 섹터'; H['Q4']='한글'
for c in (H['P4'],H['Q4']): c.font=f_hd; c.fill=PatternFill('solid',fgColor='5E6E78'); c.alignment=C
for i,(e,k) in enumerate(SK):
    H.cell(5+i,16,e).font=f_s; H.cell(5+i,17,k).font=f_s
H.column_dimensions['P'].width=22; H.column_dimensions['Q'].width=12
HR=lambda col: f'${col}${hs+1}:${col}${HL}'
# TOP10 + 기타 (A~E)
for j,tl in enumerate(['순위','티커','종목','섹터','비중'],1):
    c=H.cell(hs,j,tl); c.font=f_hd; c.fill=hd_fill; c.alignment=C
for j,w in zip(range(1,6),[6,8,22,13,10]): H.column_dimensions[L(j)].width=w
for k in range(1,11):
    rr=hs+k
    H.cell(rr,1,k).alignment=Alignment(horizontal='center')
    H.cell(rr,5,f'=LARGE({HR("I")},A{rr})').number_format='0.00%'
    H.cell(rr,2,f'=INDEX({HR("H")},MATCH(E{rr},{HR("I")},0))')
    H.cell(rr,3,f'=IF(INDEX({HR("L")},MATCH(E{rr},{HR("I")},0))="",INDEX({HR("J")},MATCH(E{rr},{HR("I")},0)),INDEX({HR("L")},MATCH(E{rr},{HR("I")},0)))')
    H.cell(rr,4,f'=INDEX({HR("M")},MATCH(E{rr},{HR("I")},0))')
    for col in range(1,6): H.cell(rr,col).font=f_fx; H.cell(rr,col).border=bd
r11=hs+11
H.cell(r11,3,'기타 (나머지 종목)').font=f_n
H.cell(r11,5,f'=SUM({HR("I")})-SUM(E{hs+1}:E{hs+10})').number_format='0.00%'
H.cell(r11+1,3,'합계').font=f_b
H.cell(r11+1,5,f'=SUM(E{hs+1}:E{r11})').number_format='0.00%'; H.cell(r11+1,5).font=f_b
H.cell(r11+2,3,'TOP10 합계').font=f_n
H.cell(r11+2,5,f'=SUM(E{hs+1}:E{hs+10})').number_format='0.00%'
H.cell(r11+3,3,'보유 종목 수(현금·선물 제외)').font=f_n
H.cell(r11+3,5,f'=COUNTIFS({HR("H")},"<>",{HR("M")},"<>현금·선물")-COUNTBLANK({HR("H")})+COUNTBLANK({HR("H")})').number_format='0'
H.cell(r11+3,5).value=f'=SUMPRODUCT(({HR("H")}<>"")*({HR("M")}<>"현금·선물"))'
# 섹터 비중 (A~E 아래)
sr0=r11+6
H.cell(sr0-1,1,'섹터 비중').font=f_b
for j,tl in enumerate(['','','섹터','','비중'],1):
    if tl: c=H.cell(sr0,j,tl); c.font=f_hd; c.fill=hd_fill; c.alignment=C
for i,(e,k) in enumerate(SK[:9]+[('','현금·선물')]):
    rr=sr0+1+i
    H.cell(rr,3,k).font=f_n
    H.cell(rr,5,f'=SUMIFS({HR("I")},{HR("M")},C{rr})').number_format='0.00%'
    H.cell(rr,5).font=f_fx; H.cell(rr,3).border=bd; H.cell(rr,5).border=bd
H.cell(sr0+12,1,'출처: Schwab Asset Management, SCHD 보유종목 전체 목록(Export All Holdings). 비중은 순자산 대비 %. 종목 정보는 투자 권유가 아닙니다.').font=f_s
# 도넛 차트
dc=DoughnutChart(); dc.title='SCHD TOP10 + 기타 (합계 100%)'; dc.holeSize=55; dc.height=10; dc.width=14
dc.add_data(Reference(H,min_col=5,min_row=hs,max_row=r11),titles_from_data=True)
dc.set_categories(Reference(H,min_col=3,min_row=hs+1,max_row=r11))
PAL=['A8661B','C98A3A','E0B36A','2F5D8A','4F7FAF','86A9CC','17784A','4E9E72','8CC3A2','7A5C8E','C9D1D6']
ser=dc.series[0]
for idx,col in enumerate(PAL):
    pt=DataPoint(idx=idx); pt.graphicalProperties=GraphicalProperties(solidFill=col); ser.dPt.append(pt)
ser.dLbls=DataLabelList(); ser.dLbls.showVal=True; ser.dLbls.numFmt='0.0%'
for k in ('showSerName','showCatName','showLegendKey','showPercent'): setattr(ser.dLbls,k,False)
dc.legend.position='r'
H.add_chart(dc,'S4')

wb._sheets=[wb[n] for n in ['사용법','최신분기_요약','SCHD_배당','연도별','분기표','국내SCHD형_분배','국내SCHD형_비교','구성종목','차트데이터']]
wb.calculation.fullCalcOnLoad=True
out=ROOT/'downloads'/'schd_dividend_tracker.xlsx'
wb.save(out); print(out)
